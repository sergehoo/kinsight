"""Rendus du rapport Capital Humain : Excel (détail + synthèse), PDF (rapport
exécutif), PowerPoint (présentation Direction).

Les trois rendent `rapport_rh.assembler(...)` tel quel : aucun calcul ici, aucune
valeur par défaut. Une case non mesurée reste VIDE dans les séries (un graphe y
montre un trou) et « N/D » dans les tableaux, avec le motif. Les graphiques sont
natifs (Excel, PowerPoint) ou vectoriels (PDF) et tracent les mêmes points que
l'écran.
"""

from __future__ import annotations

import io
from datetime import date

from .rapport_rh import comparaison_texte, libelle_filtres, valeur_texte

ETATS = {"critique": "Critique", "surveillance": "Surveillance", "normal": "Normal", "non_mesure": "Non mesuré",
         "info": "À noter"}
NOTE = ("Aucun chiffre n'est inventé : N/D signifie non mesuré, jamais zéro. Sources : Kaydan Shield (présence, "
        "effectifs déclarés, sites) et Odoo RH (référentiel), historisés chaque jour dans K-Insight.")


def _jour(iso: str) -> str:
    d = date.fromisoformat(iso)
    return d.strftime("%d/%m")


def _seuils(s: dict) -> str:
    seuils = s.get("seuils")
    if not seuils:
        return ""
    op = "<" if seuils["sens"] == "bas" else ">"
    return f"surveillance {op} {seuils['surveillance']:g} % · critique {op} {seuils['critique']:g} %"


def _valeur_signal(s: dict) -> str:
    v = s.get("valeur")
    if v is None:
        return "N/D"
    if s["cle"] == "effectif_critique":
        return f"baisse {v:.1f} %".replace(".", ",") if v >= 0 else f"hausse {-v:.1f} %".replace(".", ",")
    return f"{v:.1f} %".replace(".", ",")


# ── Excel ────────────────────────────────────────────────────────────────────


def rapport_xlsx(r: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill

    gras = Font(bold=True)
    entete = PatternFill("solid", fgColor="E8EEF9")
    wb = Workbook()

    ws = wb.active
    ws.title = "Synthèse"
    ws.append([r["titre"]])
    ws["A1"].font = Font(bold=True, size=14)
    for cle, valeur in (("Généré le", r["genere_le"]), ("Par", f"{r['auteur']} ({r['role']})"),
                        ("Périmètre", r["perimetre"]), ("Filtres", libelle_filtres(r))):
        ws.append([cle, valeur])
    ws.append([])

    def tableau(titre: str, colonnes: list[str], lignes: list[list]):
        ws.append([titre])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        ws.append(colonnes)
        for c in ws[ws.max_row]:
            c.font, c.fill = gras, entete
        for ligne in lignes:
            ws.append(ligne)
        ws.append([])

    tableau("Indicateurs de la période", ["Indicateur", "Valeur", "Unité", "Comparaison", "Source",
                                          "Jours mesurés", "Motif si non mesuré"],
            [[k["libelle"], k["valeur"] if k["valeur"] is not None else "N/D", k["unite"], comparaison_texte(k),
              k["source"], (f"{k['couverture']['jours_mesures']}/{k['couverture']['jours']}" if k.get("couverture") else ""),
              k.get("motif") or ""] for k in r["kpis"]])
    tableau("Signaux", ["Signal", "État", "Valeur", "Seuils", "Source", "Motif"],
            [[s["libelle"], ETATS[s["etat"]], _valeur_signal(s), _seuils(s), s.get("source") or "",
              s.get("motif") or ""] for s in r["signaux"]])
    tableau("Constats prioritaires", ["Sévérité", "Constat", "Détail", "Action recommandée", "Source", "Confiance"],
            [[ETATS[c["severite"]], c["titre"], c["constat"], c.get("action") or "", c.get("source") or "",
              c.get("confiance") or ""] for c in r["constats"]] or [["—", "Aucun constat : rien à signaler sur la période."]])
    score = r["score"]
    tableau("Score de gouvernance RH", ["Dimension", "Poids", "Score", "Motif"],
            [["Global", 100, score["global"] if score["global"] is not None else "N/D", score.get("motif_global") or ""]]
            + [[d["label"], d["weight"], d["score"] if d["score"] is not None else "N/D", d.get("motif") or ""]
               for d in score["dimensions"]])
    tableau("Sources et fraîcheur", ["Source", "État", "Dernière collecte"],
            [[s["source"], s["etat"], s["derniere_collecte"] or "jamais"] for s in r["sources"]])
    ws.append([NOTE])
    for colonne, largeur in zip("ABCDEFG", (38, 16, 12, 22, 18, 14, 60)):
        ws.column_dimensions[colonne].width = largeur
    for ligne in ws.iter_rows():
        for c in ligne:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (3, 4, 7) and isinstance(c.value, str))

    # Série quotidienne : une ligne par jour, une case VIDE pour un jour non mesuré.
    if r["series"]:
        wq = wb.create_sheet("Série quotidienne")
        cles = list(r["series"])
        wq.append(["Date"] + [r["series"][c]["libelle"] for c in cles])
        for c in wq[1]:
            c.font, c.fill = gras, entete
        jours = [p["date"] for p in r["series"][cles[0]]["points"]]
        for i, jour in enumerate(jours):
            wq.append([date.fromisoformat(jour)] + [r["series"][c]["points"][i]["value"] for c in cles])
        wq.column_dimensions["A"].width = 12
        if "hr.taux_presence_jour" in r["series"] and len(jours) > 1:
            graphe = LineChart()
            graphe.title, graphe.y_axis.title, graphe.height, graphe.width = "Taux de présence", "%", 8, 18
            col = cles.index("hr.taux_presence_jour") + 2
            graphe.add_data(Reference(wq, min_col=col, min_row=1, max_row=len(jours) + 1), titles_from_data=True)
            graphe.set_categories(Reference(wq, min_col=1, min_row=2, max_row=len(jours) + 1))
            graphe.display_blanks = "gap"   # un jour non mesuré reste un trou
            wq.add_chart(graphe, "G2")

    structure = r.get("structure")
    if structure and structure.get("repartitions"):
        wst = wb.create_sheet("Structure (Shield)")
        wst.append(["Notion", "Catégorie", "Effectif", "Part (%)"])
        for c in wst[1]:
            c.font, c.fill = gras, entete
        for notion, lignes in structure["repartitions"].items():
            for l in lignes:
                wst.append([structure.get("libelles", {}).get(notion, notion), l["libelle"], l["valeur"], l["part"]])
        if structure.get("contrats"):
            wst.append([])
            for h, n in structure["contrats"]["par_horizon"].items():
                wst.append(["Contrats à échéance", f"≤ {h} jours", n, None])
        wst.column_dimensions["A"].width, wst.column_dimensions["B"].width = 22, 34

    if r["sites"]:
        wsite = wb.create_sheet("Sites")
        wsite.append(["Site", "Taux de présence (%)", "Jours mesurés"])
        for c in wsite[1]:
            c.font, c.fill = gras, entete
        for l in r["sites"]:
            wsite.append([l["libelle"], l["value"], f"{l['couverture']['jours_mesures']}/{l['couverture']['jours']}"])
        wsite.column_dimensions["A"].width = 32
        barres = BarChart()
        barres.type, barres.title, barres.height, barres.width = "bar", "Présence par site", 8, 16
        barres.add_data(Reference(wsite, min_col=2, min_row=1, max_row=len(r["sites"]) + 1), titles_from_data=True)
        barres.set_categories(Reference(wsite, min_col=1, min_row=2, max_row=len(r["sites"]) + 1))
        wsite.add_chart(barres, "E2")

    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────


def _graphe_pdf(points: list[dict], largeur: float, hauteur: float):
    """Courbe de présence ; un jour non mesuré coupe la ligne (segments séparés)."""
    from reportlab.graphics.charts.lineplots import LinePlot
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors

    segments, courant = [], []
    for i, p in enumerate(points):
        if p["value"] is None:
            if courant:
                segments.append(courant)
            courant = []
        else:
            courant.append((i, p["value"]))
    if courant:
        segments.append(courant)
    dessin = Drawing(largeur, hauteur)
    if not segments:
        dessin.add(String(10, hauteur / 2, "Aucune journée mesurée sur la période.", fontSize=9))
        return dessin
    lp = LinePlot()
    lp.x, lp.y, lp.width, lp.height = 30, 20, largeur - 45, hauteur - 30
    lp.data = segments
    for i in range(len(segments)):
        lp.lines[i].strokeColor = colors.HexColor("#2D74E0")
        lp.lines[i].strokeWidth = 1.6
    valeurs = [v for seg in segments for _, v in seg]
    lp.yValueAxis.valueMin = max(0, min(valeurs) - 5)
    lp.yValueAxis.valueMax = min(100, max(valeurs) + 5)
    lp.xValueAxis.valueMin, lp.xValueAxis.valueMax = 0, max(1, len(points) - 1)
    pas = max(1, len(points) // 8)
    lp.xValueAxis.valueSteps = list(range(0, len(points), pas))
    lp.xValueAxis.labelTextFormat = lambda i: _jour(points[int(i)]["date"]) if 0 <= int(i) < len(points) else ""
    lp.xValueAxis.labels.fontSize = lp.yValueAxis.labels.fontSize = 7
    dessin.add(lp)
    return dessin


def rapport_pdf(r: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    petit = styles["BodyText"].clone("petit", fontSize=8, leading=10)
    sortie = io.BytesIO()
    doc = SimpleDocTemplate(sortie, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                            topMargin=1.4 * cm, bottomMargin=1.4 * cm, title=r["titre"])
    largeur = A4[0] - 3.2 * cm

    def table(lignes, largeurs):
        t = Table([[Paragraph(str(c), petit) for c in l] for l in lignes], colWidths=largeurs, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF9")),
                               ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D1CE")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    h = []
    h.append(Paragraph(r["titre"], styles["Title"]))
    h.append(Paragraph(f"{libelle_filtres(r)}<br/>Périmètre : {r['perimetre']} · généré le {r['genere_le'][:16].replace('T', ' ')} "
                       f"par {r['auteur']} ({r['role']})", petit))
    h.append(Spacer(1, 8))

    h.append(Paragraph("Synthèse", styles["Heading2"]))
    for phrase in r["synthese"]["resume"]:
        h.append(Paragraph(f"• {phrase}", styles["BodyText"]))
    h.append(Spacer(1, 6))

    h.append(Paragraph("Indicateurs de la période", styles["Heading2"]))
    h.append(table([["Indicateur", "Valeur", "Comparaison", "Source", "Motif si non mesuré"]]
                   + [[k["libelle"], valeur_texte(k), comparaison_texte(k), k["source"], k.get("motif") or ""]
                      for k in r["kpis"]],
                   [4.6 * cm, 2.2 * cm, 3.0 * cm, 2.6 * cm, largeur - 12.4 * cm]))

    presence = r["series"].get("hr.taux_presence_jour")
    if presence:
        # Le titre ne se sépare pas de son graphe en bas de page.
        h.append(KeepTogether([Paragraph("Taux de présence, jour par jour", styles["Heading2"]),
                               _graphe_pdf(presence["points"], largeur, 5 * cm)]))

    structure = r.get("structure")
    if structure and structure.get("repartitions", {}).get("contrat"):
        h.append(Paragraph("Effectifs par type de contrat (Kaydan Shield, aujourd'hui)", styles["Heading2"]))
        h.append(table([["Contrat", "Effectif", "Part"]]
                       + [[l["libelle"], l["valeur"], f"{l['part']:g} %" if l["part"] is not None else ""]
                          for l in structure["repartitions"]["contrat"]],
                       [6 * cm, 3 * cm, largeur - 9 * cm]))

    h.append(Paragraph("Signaux", styles["Heading2"]))
    h.append(table([["Signal", "État", "Valeur", "Seuils", "Motif"]]
                   + [[s["libelle"], ETATS[s["etat"]], _valeur_signal(s), _seuils(s), s.get("motif") or ""]
                      for s in r["signaux"]],
                   [4.4 * cm, 2.2 * cm, 2.4 * cm, 4.2 * cm, largeur - 13.2 * cm]))

    h.append(Paragraph("Constats et recommandations", styles["Heading2"]))
    if r["constats"]:
        for c in r["constats"]:
            h.append(Paragraph(f"<b>[{ETATS[c['severite']]}] {c['titre']}</b> — {c['constat']}", styles["BodyText"]))
            if c.get("action"):
                h.append(Paragraph(f"<i>Action recommandée :</i> {c['action']}", petit))
            h.append(Spacer(1, 3))
    else:
        h.append(Paragraph("Aucun constat : aucun seuil franchi, aucun écart notable sur la période.", styles["BodyText"]))

    score = r["score"]
    h.append(Paragraph("Score de gouvernance RH", styles["Heading2"]))
    h.append(Paragraph(("Global : " + (f"{score['global']:g}/100" if score["global"] is not None else "N/D")
                        + (f" — {score['motif_global']}" if score.get("motif_global") else "")), styles["BodyText"]))

    h.append(Paragraph("Sources et fraîcheur", styles["Heading2"]))
    h.append(table([["Source", "État", "Dernière collecte"]]
                   + [[s["source"], s["etat"], s["derniere_collecte"] or "jamais"] for s in r["sources"]],
                   [5 * cm, 4 * cm, largeur - 9 * cm]))
    h.append(Spacer(1, 8))
    h.append(Paragraph(NOTE, petit))
    doc.build(h)
    return sortie.getvalue()


# ── PowerPoint ───────────────────────────────────────────────────────────────


def rapport_pptx(r: dict) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    prs = Presentation()
    largeur_origine = prs.slide_width
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    titre_seul, titre_contenu, vide = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[5]
    facteur = prs.slide_width / largeur_origine

    def etirer(diapo):
        """Le gabarit par défaut est en 4:3 : ses zones restaient tassées à gauche
        d'une diapositive 16:9. On les élargit dans la même proportion."""
        for zone in diapo.placeholders:
            # Les quatre à la fois : une zone héritée du gabarit n'a pas de
            # position propre, et n'en fixer qu'une partie remettait le reste à 0.
            gauche, haut, larg, haut_ = zone.left, zone.top, zone.width, zone.height
            zone.left, zone.top = int(gauche * facteur), haut
            zone.width, zone.height = int(larg * facteur), haut_
        return diapo

    def puces(titre: str, lignes: list[str], vide_texte: str):
        s = etirer(prs.slides.add_slide(titre_contenu))
        s.shapes.title.text = titre
        corps = s.placeholders[1].text_frame
        corps.clear()
        for i, ligne in enumerate(lignes or [vide_texte]):
            p = corps.paragraphs[0] if i == 0 else corps.add_paragraph()
            p.text = ligne
            p.font.size = Pt(16 if len(lignes) <= 5 else 13)
        return s

    s = etirer(prs.slides.add_slide(titre_seul))
    s.shapes.title.text = r["titre"]
    s.placeholders[1].text = (f"Périmètre : {r['perimetre']}\n{libelle_filtres(r)}\n"
                              f"Généré le {r['genere_le'][:10]} par {r['auteur']}")

    puces("Synthèse", r["synthese"]["resume"], "Aucun seuil franchi sur la période : rien à signaler.")

    s = etirer(prs.slides.add_slide(vide))
    s.shapes.title.text = "Indicateurs clés"
    lignes = [k for k in r["kpis"]][:12]
    tableau = s.shapes.add_table(len(lignes) + 1, 4, Inches(0.5), Inches(1.4), Inches(12.3), Inches(5.6)).table
    for j, entete in enumerate(("Indicateur", "Valeur", "Comparaison", "Source")):
        tableau.cell(0, j).text = entete
    for i, k in enumerate(lignes, start=1):
        for j, texte in enumerate((k["libelle"], valeur_texte(k), comparaison_texte(k), k["source"])):
            tableau.cell(i, j).text = texte
    for ligne in tableau.rows:
        for cellule in ligne.cells:
            for p in cellule.text_frame.paragraphs:
                p.font.size = Pt(12)

    s = etirer(prs.slides.add_slide(vide))
    s.shapes.title.text = "Tendance — taux de présence"
    presence = r["series"].get("hr.taux_presence_jour")
    if presence and any(p["value"] is not None for p in presence["points"]):
        donnees = CategoryChartData()
        donnees.categories = [_jour(p["date"]) for p in presence["points"]]
        donnees.add_series("Taux de présence (%)", [p["value"] for p in presence["points"]])  # None = trou
        graphe = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(0.6), Inches(1.4), Inches(12.1), Inches(5.6), donnees).chart
        graphe.has_legend = True
        graphe.legend.position, graphe.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        # Un taux est borné à 100 : l'échelle automatique montait à 200 %.
        mesures = [p["value"] for p in presence["points"] if p["value"] is not None]
        axe = graphe.value_axis
        axe.maximum_scale = 100
        axe.minimum_scale = max(0, (int(min(mesures)) // 10) * 10 - 10)
        axe.major_unit = 5 if 100 - axe.minimum_scale <= 30 else 10
        axe.tick_labels.font.size = Pt(10)
        graphe.category_axis.tick_labels.font.size = Pt(9)
    else:
        boite = s.shapes.add_textbox(Inches(0.6), Inches(2), Inches(12), Inches(1)).text_frame
        boite.text = "Aucune journée mesurée sur la période (ou périmètre non rattaché aux sites Shield)."

    risques = [f"[{ETATS[x['etat']]}] {x['libelle']} : {_valeur_signal(x)}" for x in r["signaux"]
               if x["etat"] in ("critique", "surveillance")]
    risques += [f"Non mesuré — {x['libelle']}" for x in r["signaux"] if x["etat"] == "non_mesure"][:3]
    puces("Risques et signaux", risques, "Aucun signal en surveillance ni critique.")

    recommandations = [f"{c['titre']} : {c['action']}" for c in r["constats"] if c.get("action")]
    puces("Recommandations", recommandations, "Aucune action recommandée : aucun constat sur la période.")

    puces("Sources et méthode",
          [f"{x['source']} — {x['etat']}, dernière collecte {x['derniere_collecte'] or 'jamais'}"
           for x in r["sources"]] + [NOTE], "")

    sortie = io.BytesIO()
    prs.save(sortie)
    return sortie.getvalue()


RENDUS = {
    "xlsx": (rapport_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf": (rapport_pdf, "application/pdf"),
    "pptx": (rapport_pptx, "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
}
