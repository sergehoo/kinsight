"""Exports du rapport Capital Humain (lot 9).

Ce qui est protégé : les trois formats rendent les chiffres de l'écran, sous les
mêmes filtres et le même périmètre ; un N/D reste N/D (jamais 0) ; un périmètre
restreint n'exporte pas les chiffres Groupe ; un filtre hors périmètre est
refusé ; chaque extraction est tracée (qui, rôle, périmètre, filtres, format).
"""

from __future__ import annotations

import io
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.core.cache import cache
from rest_framework.test import APITestCase

from apps.audit.models import AccessLog
from apps.organizations.models import Subsidiary

from .tests_decision import _jours, _site
from .tests_historique import HIER, _poser

User = get_user_model()
URL = "/api/v1/governance/hr/rapport.{}?periode=7j"


class RapportExportTest(APITestCase):
    def setUp(self):
        cache.clear()
        self.kre = Subsidiary.objects.create(code="KRE", name="Kaydan Real Estate")
        self.drh = User.objects.create_user("drh.r", password=None, role="DRH", is_group_scope=True)
        self.drh_kre = User.objects.create_user("drh.r.kre", password=None, role="DRH", is_group_scope=False)
        self.drh_kre.subsidiaries.add(self.kre)
        self.lecteur = User.objects.create_user("lect.r", password=None, role="READER", is_group_scope=True)
        _jours(0, 6, 80, 20, retards=4)
        _poser("hr.presents_jour", HIER - timedelta(days=6), None, statut="unknown")   # un trou
        _site("1", "Chantier A", 60, 40)
        _site("2", "Siège", 95, 5)

    def _get(self, user, ext, extra=""):
        self.client.force_authenticate(user)
        return self.client.get(URL.format(ext) + extra)

    def _classeur(self, user, extra=""):
        from openpyxl import load_workbook

        r = self._get(user, "xlsx", extra)
        self.assertEqual(r.status_code, 200, r.content[:300])
        return load_workbook(io.BytesIO(r.content))

    @staticmethod
    def _ligne(ws, libelle):
        for ligne in ws.iter_rows(values_only=True):
            if ligne and ligne[0] == libelle:
                return ligne
        raise AssertionError(f"ligne « {libelle} » absente")

    def test_excel_reprend_les_chiffres_et_les_n_d(self):
        wb = self._classeur(self.drh)
        ws = wb["Synthèse"]
        self.assertEqual(self._ligne(ws, "Taux de présence")[1], 80.0)
        turnover = self._ligne(ws, "Turnover")
        self.assertEqual(turnover[1], "N/D", "un indicateur sans source n'est jamais 0")
        self.assertTrue(turnover[6])
        self.assertIn("Période : 7 derniers jours", self._ligne(ws, "Filtres")[1])
        # Série quotidienne : le jour non mesuré est une case VIDE.
        serie = list(wb["Série quotidienne"].iter_rows(values_only=True))
        self.assertEqual(serie[0][0], "Date")
        self.assertIn(None, [l[1] for l in serie[1:]])
        self.assertTrue(wb["Série quotidienne"]._charts, "graphique natif présent")
        self.assertEqual([l[0] for l in list(wb["Sites"].iter_rows(values_only=True))[1:]], ["Siège", "Chantier A"])

    def test_les_filtres_s_appliquent(self):
        wb = self._classeur(self.drh, "&site=1")
        self.assertEqual(self._ligne(wb["Synthèse"], "Taux de présence")[1], 60.0)
        self.assertNotIn("Sites", wb.sheetnames, "pas de classement des sites sous un filtre site")

    def test_perimetre_restreint_sans_chiffres_groupe(self):
        wb = self._classeur(self.drh_kre)
        ligne = self._ligne(wb["Synthèse"], "Taux de présence")
        self.assertEqual(ligne[1], "N/D")
        self.assertIn("Non rattaché", ligne[6])
        self.assertNotIn("Sites", wb.sheetnames)
        self.assertEqual(self._ligne(wb["Synthèse"], "Périmètre")[1], "KRE")

    def test_filtre_hors_perimetre_refuse(self):
        Subsidiary.objects.create(code="KSH", name="Kaydan Shield")
        self.assertEqual(self._get(self.drh_kre, "pdf", "&subsidiary=KSH").status_code, 403)

    def test_pdf_et_pptx(self):
        from pptx import Presentation

        pdf = self._get(self.drh, "pdf")
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertIn("attachment", pdf["Content-Disposition"])

        r = self._get(self.drh, "pptx")
        self.assertEqual(r.status_code, 200)
        prs = Presentation(io.BytesIO(r.content))
        titres = [s.shapes.title.text for s in prs.slides]
        self.assertEqual(titres, ["Rapport Capital Humain", "Synthèse", "Indicateurs clés",
                                  "Tendance — taux de présence", "Risques et signaux", "Recommandations",
                                  "Sources et méthode"])
        graphes = [sh for sh in prs.slides[3].shapes if sh.has_chart]
        self.assertEqual(len(graphes), 1)
        valeurs = list(graphes[0].chart.plots[0].series[0].values)
        self.assertIn(None, valeurs, "le jour non mesuré reste un trou dans la courbe")

    def test_format_inconnu_et_droits(self):
        self.assertEqual(self._get(self.drh, "docx").status_code, 400)
        self.assertEqual(self._get(self.lecteur, "pdf").status_code, 403)

    def test_extraction_tracee(self):
        self._get(self.drh, "pptx", "&site=1")
        trace = AccessLog.objects.filter(action="export_hr_report").last()
        self.assertEqual(trace.payload["format"], "pptx")
        self.assertEqual(trace.payload["role"], "DRH")
        self.assertEqual(trace.payload["site"], "1")
        self.assertEqual(trace.subsidiary_scope, ["*"])
        self.assertEqual(trace.user_role, "DRH")
