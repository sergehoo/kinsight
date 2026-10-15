"""Rapports programmés (lot 9b).

Ce qui est protégé : un destinataire reçoit le rapport sous SES droits (pas ceux
de l'auteur) ; on ne programme pas vers une adresse libre ni vers un utilisateur
sans accès RH ; sans SMTP rien ne part et le journal le dit ; une échéance ne
part qu'une fois même si la tâche est rejouée.
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta

from django.contrib.auth import get_user_model
from django.core import mail
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.test import APITestCase

from apps.organizations.models import Subsidiary

from . import rapports_programmes as P
from .models import EnvoiRapport, RapportProgramme
from .tests_decision import _jours

User = get_user_model()
URL = "/api/v1/governance/hr/rapports-programmes/"
SMTP = override_settings(EMAIL_HOST="smtp.test", EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")


def _a(annee, mois, jour, heure=9):
    return timezone.make_aware(datetime(annee, mois, jour, heure, 0))


class EcheancesTest(TestCase):
    def test_quotidien(self):
        self.assertEqual(P.prochaine_echeance("quotidien", _a(2026, 10, 15, 6)), _a(2026, 10, 15, 7))
        self.assertEqual(P.prochaine_echeance("quotidien", _a(2026, 10, 15, 9)), _a(2026, 10, 16, 7))

    def test_hebdomadaire_le_lundi(self):
        self.assertEqual(P.prochaine_echeance("hebdomadaire", _a(2026, 10, 15)), _a(2026, 10, 19, 7))   # jeudi → lundi
        self.assertEqual(P.prochaine_echeance("hebdomadaire", _a(2026, 10, 19, 8)), _a(2026, 10, 26, 7))

    def test_mensuel_et_changement_d_annee(self):
        self.assertEqual(P.prochaine_echeance("mensuel", _a(2026, 10, 15)), _a(2026, 11, 1, 7))
        self.assertEqual(P.prochaine_echeance("mensuel", _a(2026, 12, 3)), _a(2027, 1, 1, 7))
        self.assertEqual(P.prochaine_echeance("mensuel", _a(2026, 11, 1, 6)), _a(2026, 11, 1, 7))


class ProgrammationTest(APITestCase):
    def setUp(self):
        cache.clear()
        self.kre = Subsidiary.objects.create(code="KRE", name="Kaydan Real Estate")
        Subsidiary.objects.create(code="KSH", name="Kaydan Shield")
        self.drh = User.objects.create_user("drh.p", password=None, role="DRH", is_group_scope=True,
                                            email="drh@kaydan.test")
        self.drh_kre = User.objects.create_user("drh.p.kre", password=None, role="DRH", is_group_scope=False,
                                                email="kre@kaydan.test")
        self.drh_kre.subsidiaries.add(self.kre)
        self.lecteur = User.objects.create_user("lect.p", password=None, role="READER", is_group_scope=True,
                                                email="lect@kaydan.test")
        _jours(0, 7, 80, 20)

    def _creer(self, user, **corps):
        self.client.force_authenticate(user)
        base = {"nom": "Hebdo DRH", "frequence": "hebdomadaire", "format": "xlsx", "filtres": {"periode": "7j"}}
        return self.client.post(URL, {**base, **corps}, format="json")

    def test_programmer_pour_soi_et_pour_un_rh(self):
        r = self._creer(self.drh, destinataires=[self.drh.pk, self.drh_kre.pk])
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(len(r.json()["destinataires"]), 2)
        self.assertGreater(RapportProgramme.objects.get().prochain_envoi, timezone.now())

    def test_destinataires_autorises_seulement(self):
        self.assertEqual(self._creer(self.drh, destinataires=[self.lecteur.pk]).status_code, 400,
                         "un utilisateur sans accès RH ne reçoit pas de rapport RH")
        r = self._creer(self.drh_kre, destinataires=[self.drh.pk])
        self.assertEqual(r.status_code, 400, "un RH de filiale ne programme que pour lui-même")
        self.assertEqual(self._creer(self.drh_kre).status_code, 201)

    def test_filtres_controles_a_la_programmation(self):
        self.assertEqual(self._creer(self.drh_kre, filtres={"periode": "7j", "subsidiary": "KSH"}).status_code, 403)
        self.assertEqual(self._creer(self.drh, filtres={"periode": "lune"}).status_code, 400)
        self.assertEqual(self._creer(self.drh, format="docx").status_code, 400)

    @SMTP
    def test_chaque_destinataire_recoit_sous_ses_droits(self):
        from openpyxl import load_workbook

        r = self._creer(self.drh, destinataires=[self.drh.pk, self.drh_kre.pk])
        rapport = RapportProgramme.objects.get(pk=r.json()["id"])
        P.envoyer(rapport)
        self.assertEqual(len(mail.outbox), 2)
        valeurs = {}
        for message in mail.outbox:
            nom, contenu, _ = message.attachments[0]
            self.assertTrue(nom.endswith(".xlsx"))
            ws = load_workbook(io.BytesIO(contenu))["Synthèse"]
            ligne = next(l for l in ws.iter_rows(values_only=True) if l and l[0] == "Taux de présence")
            valeurs[message.to[0]] = ligne[1]
        self.assertEqual(valeurs["drh@kaydan.test"], 80.0)
        self.assertEqual(valeurs["kre@kaydan.test"], "N/D", "jamais le chiffre Groupe sous l'étiquette filiale")
        self.assertEqual(EnvoiRapport.objects.filter(statut="envoye").count(), 2)

    def test_sans_smtp_rien_ne_part_et_le_journal_le_dit(self):
        r = self._creer(self.drh)
        P.envoyer(RapportProgramme.objects.get(pk=r.json()["id"]))
        self.assertEqual(len(mail.outbox), 0)
        trace = EnvoiRapport.objects.get()
        self.assertEqual(trace.statut, "ignore")
        self.assertIn("SMTP non configuré", trace.detail)

    @SMTP
    def test_destinataire_qui_a_perdu_l_acces(self):
        r = self._creer(self.drh, destinataires=[self.drh_kre.pk])
        self.drh_kre.role = "READER"
        self.drh_kre.save()
        P.envoyer(RapportProgramme.objects.get(pk=r.json()["id"]))
        self.assertEqual(len(mail.outbox), 0)
        self.assertEqual(EnvoiRapport.objects.get().statut, "ignore")

    @SMTP
    def test_une_echeance_ne_part_qu_une_fois(self):
        self._creer(self.drh)
        RapportProgramme.objects.update(prochain_envoi=timezone.now() - timedelta(minutes=1))
        maintenant = timezone.now()
        self.assertEqual(P.envoyer_les_dus(maintenant)["envoyes"], 1)
        self.assertEqual(P.envoyer_les_dus(maintenant)["rapports"], 0, "tâche rejouée : pas de second envoi")
        self.assertGreater(RapportProgramme.objects.get().prochain_envoi, maintenant)

    def test_suspension_et_isolation(self):
        r = self._creer(self.drh)
        pk = r.json()["id"]
        self.client.force_authenticate(self.drh)
        self.assertFalse(self.client.patch(f"{URL}{pk}/", {"actif": False}, format="json").json()["actif"])
        # Un RH de filiale ne voit ni ne déclenche la programmation d'un autre.
        self.client.force_authenticate(self.drh_kre)
        self.assertEqual(self.client.get(URL).json()["rapports"], [])
        self.assertEqual(self.client.post(f"{URL}{pk}/envoyer/").status_code, 404)
        self.client.force_authenticate(self.lecteur)
        self.assertEqual(self.client.get(URL).status_code, 403)
