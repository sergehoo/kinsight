"""Profils RBAC du Capital Humain (lot 11) — la matrice, côté serveur.

DG Groupe : tout. DRH Groupe : tout le Capital Humain, paie / carrière /
performance comprises. Directeur de filiale : tous les métiers, TOUJOURS limité à
ses filiales. RH filiale : le Capital Humain de sa filiale, sans les volets
renforcés. Assistant RH : les seules pages affectées. API, exports et lectures
d'historique obéissent aux mêmes règles ; aucun filtre n'élargit un périmètre.
"""

from __future__ import annotations

import io

from django.contrib.auth import get_user_model
from django.core.cache import cache
from rest_framework.test import APITestCase

from apps.accounts.rbac import permissions_for
from apps.organizations.models import Subsidiary

from .tests_decision import _jours

User = get_user_model()


class ProfilsRhTest(APITestCase):
    def setUp(self):
        cache.clear()
        self.kre = Subsidiary.objects.create(code="KRE", name="Kaydan Real Estate")
        Subsidiary.objects.create(code="KSH", name="Kaydan Shield")
        mk = User.objects.create_user
        self.dg = mk("dg", password=None, role="DG_GROUP", is_group_scope=True)
        self.drh = mk("drh.g", password=None, role="DRH", is_group_scope=True)
        # Le drapeau « groupe » coché par erreur ne doit rien ouvrir à un directeur de filiale.
        self.dir_kre = mk("dir.kre", password=None, role="DIR_FILIALE", is_group_scope=True)
        self.rh_kre = mk("rh.kre", password=None, role="DRH", is_group_scope=False)
        self.assistant = mk("assistant", password=None, role="ASSISTANT_RH", is_group_scope=False,
                            affectations_rh=["hr-attendance"])
        self.lecteur = mk("lecteur", password=None, role="READER", is_group_scope=True)
        for u in (self.dir_kre, self.rh_kre, self.assistant):
            u.subsidiaries.add(self.kre)
        _jours(0, 7, 80, 20)

    def _statut(self, user, url):
        self.client.force_authenticate(user)
        return self.client.get(url).status_code

    def test_permissions_par_profil(self):
        renforcees = {"view_hr_payroll", "view_hr_career", "view_hr_performance"}
        self.assertTrue(renforcees <= set(permissions_for(self.drh)))
        self.assertFalse(renforcees & set(permissions_for(self.rh_kre)))
        self.assertIn("view_hr_payroll", permissions_for(self.dir_kre))
        self.assertNotIn("view_hr_performance", permissions_for(self.dir_kre))
        self.assertIn("hr_module:hr-attendance", permissions_for(self.assistant))

    def test_directeur_de_filiale_toujours_restreint(self):
        self.assertFalse(self.dir_kre.scope().is_group)
        self.assertEqual(set(self.dir_kre.scope().subsidiaries), {"KRE"})

    def test_pages_renforcees(self):
        attendus = {
            "/api/v1/governance/module/hr-payroll/": {"dg": 200, "drh": 200, "dir_kre": 200, "rh_kre": 403, "assistant": 403},
            "/api/v1/governance/module/hr-performance/": {"dg": 200, "drh": 200, "dir_kre": 403, "rh_kre": 403},
            "/api/v1/governance/module/hr-career/": {"drh": 200, "rh_kre": 403},
        }
        for url, par_profil in attendus.items():
            for profil, code in par_profil.items():
                with self.subTest(url=url, profil=profil):
                    self.assertEqual(self._statut(getattr(self, profil), url), code)

    def test_assistant_rh_pages_affectees_seulement(self):
        self.assertEqual(self._statut(self.assistant, "/api/v1/governance/module/hr-attendance/"), 200)
        for url in ("/api/v1/governance/module/hr-executive/", "/api/v1/governance/hr/signaux/",
                    "/api/v1/governance/hr/decision/", "/api/v1/governance/hr/score/",
                    "/api/v1/governance/hr/rapport.pdf", "/api/v1/governance/hr/rapports-programmes/"):
            with self.subTest(url=url):
                self.assertEqual(self._statut(self.assistant, url), 403)

    def test_masse_salariale_reservee_a_la_paie(self):
        url = "/api/v1/governance/hr/history/?metric=hr.masse_salariale"
        self.assertEqual(self._statut(self.rh_kre, url), 403)
        self.assertEqual(self._statut(self.drh, url), 200)

    def test_export_d_un_rh_de_filiale(self):
        from openpyxl import load_workbook

        self.client.force_authenticate(self.rh_kre)
        r = self.client.get("/api/v1/governance/hr/rapport.xlsx?periode=7j")
        self.assertEqual(r.status_code, 200)
        lignes = {l[0]: l for l in load_workbook(io.BytesIO(r.content))["Synthèse"].iter_rows(values_only=True) if l}
        self.assertEqual(lignes["Masse salariale"][6], "Réservé aux profils habilités à la paie.")
        self.assertEqual(lignes["Taux de présence"][1], "N/D", "chiffre Groupe jamais servi à une filiale")

    def test_aucun_filtre_n_elargit_le_perimetre(self):
        for profil in (self.dir_kre, self.rh_kre):
            for url in ("/api/v1/governance/hr/history/?metric=hr.employes_actifs_odoo&subsidiary=KSH",
                        "/api/v1/governance/hr/signaux/?subsidiary=KSH",
                        "/api/v1/governance/hr/rapport.xlsx?subsidiary=KSH",
                        "/api/v1/governance/module/hr-executive/?subsidiary=KSH"):
                with self.subTest(profil=profil.username, url=url):
                    self.assertEqual(self._statut(profil, url), 403)

    def test_directeur_de_filiale_ne_voit_pas_le_groupe(self):
        self.client.force_authenticate(self.dir_kre)
        d = self.client.get("/api/v1/governance/hr/history/?metric=hr.taux_presence_jour&periode=7j").json()
        self.assertEqual(d["status"], "disconnected")
        self.assertIn("Non rattaché", d["reason"])

    def test_lecteur_sans_capital_humain(self):
        for url in ("/api/v1/governance/module/hr-attendance/", "/api/v1/governance/hr/signaux/",
                    "/api/v1/governance/hr/rapport.pdf"):
            with self.subTest(url=url):
                self.assertEqual(self._statut(self.lecteur, url), 403)

    def test_profil_expose_au_frontend(self):
        self.client.force_authenticate(self.assistant)
        moi = self.client.get("/api/v1/auth/me/").json()
        self.assertIn("hr_module:hr-attendance", moi["permissions"])
        self.assertEqual(moi["role"], "ASSISTANT_RH")
