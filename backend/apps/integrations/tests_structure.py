"""Structure des effectifs Shield : champs live, filtres, périmètre, aucune donnée personnelle."""

from __future__ import annotations

from datetime import date
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import TestCase
from rest_framework.test import APITestCase

from apps.organizations.models import Subsidiary

from .models import DataConnector, DataSource, SourceStatus, SourceType
from .shield_structure import lire_structure, normaliser

User = get_user_model()
AUJ = date(2026, 10, 16)

EMPLOYES = [
    {"id": 1, "first_name": "Awa", "last_name": "Koné", "company": 3, "company_name": "Kaydan RE",
     "department": 10, "contract_type": "CDI", "status": "active", "gender": "F",
     "birth_date": "1990-05-01", "hire_date": "2020-01-15"},
    {"id": 2, "first_name": "Yao", "last_name": "N'Guessan", "company": 3, "company_name": "Kaydan RE",
     "department": 11, "contract_type": "CDD", "status": "active", "gender": "M",
     "birth_date": "2003-02-01", "hire_date": "2026-03-01"},
    {"id": 3, "first_name": "Ama", "last_name": "Diallo", "company": 4, "company_name": "Kaydan Shield",
     "department": {"id": 10, "name": "Finance"}, "contract_type": "CDD", "status": "inactive"},
]
OUVRIERS = [{"id": 9, "first_name": "Koffi", "site": 7, "site_name": "Chantier Bingerville", "status": "active",
             "trade": "Maçon"}]
DEPARTEMENTS = [{"id": 10, "name": "Finance"}, {"id": 11, "name": "Travaux"}]
CONTRATS = [{"id": 50, "employee": 2, "contract_type": "CDD", "end_date": "2026-11-05"},
            {"id": 51, "employee": 3, "contract_type": "CDD", "end_date": "2026-12-31"}]


def _charge():
    return {"status": "connected", "source": "kshieldApp", "collecte_le": "2026-10-16T08:00:00+00:00",
            **normaliser(EMPLOYES, OUVRIERS, departements=DEPARTEMENTS, contrats=CONTRATS, aujourd_hui=AUJ)}


class NormalisationTest(TestCase):
    def test_categories_live_et_aucune_donnee_personnelle(self):
        charge = _charge()
        texte = repr(charge)
        for nom in ("Awa", "Koné", "Yao", "Koffi", "1990-05-01"):
            self.assertNotIn(nom, texte, "aucun nom ni date de naissance ne quitte la collecte")
        s = lire_structure(charge, {}, AUJ)
        self.assertEqual(s["effectifs"], {"total": 4, "employes": 3, "ouvriers": 1, "actifs": 3, "inactifs": 1})
        contrats = {r["libelle"]: r["valeur"] for r in s["repartitions"]["contrat"]}
        self.assertEqual(contrats, {"CDD": 2, "CDI": 1}, "les catégories sont celles de Shield, jamais une liste en dur")
        departements = {r["libelle"]: r["valeur"] for r in s["repartitions"]["departement"]}
        self.assertEqual(departements, {"Finance": 2, "Travaux": 1}, "id et objet résolus par le référentiel")
        societes = {r["libelle"] for r in s["repartitions"]["societe"]}
        self.assertEqual(societes, {"Kaydan RE", "Kaydan Shield"})
        self.assertIn("< 25 ans", {r["libelle"] for r in s["repartitions"]["tranche_age"]})

    def test_echeances_de_contrat(self):
        s = lire_structure(_charge(), {}, AUJ)
        self.assertEqual(s["contrats"]["par_horizon"], {"30": 1, "60": 1, "90": 2})
        self.assertEqual(s["contrats"]["par_contrat"][0]["libelle"], "CDD")

    def test_filtres_depuis_les_donnees_live(self):
        s = lire_structure(_charge(), {"contrat": "CDD"}, AUJ)
        self.assertEqual(s["effectifs"]["total"], 2)
        self.assertEqual([f["libelle"] for f in s["filtres"]["contrat"]], ["CDD", "CDI"])
        self.assertEqual(s["applique"], {"contrat": "CDD"})

    def test_champ_absent_reste_nd_avec_motif(self):
        charge = {"status": "connected", **normaliser([{"id": 1, "status": "active"}], [], aujourd_hui=AUJ)}
        s = lire_structure(charge, {}, AUJ)
        self.assertNotIn("contrat", s["repartitions"])
        self.assertIn("type de contrat", s["motifs"]["contrat"])
        self.assertIsNone(s["contrats"])
        self.assertTrue(s["contrats_motif"])


class SourceIndisponibleTest(TestCase):
    def test_source_non_connectee_donne_nd_jamais_zero(self):
        s = lire_structure({"status": "disconnected", "detail": "Source kaydan-shield non configurée.", "lignes": []}, {})
        self.assertIsNone(s["effectifs"])
        self.assertEqual(s["motifs"]["contrat"], "Source kaydan-shield non configurée.")


class StructureApiTest(APITestCase):
    def setUp(self):
        cache.clear()
        self.kre = Subsidiary.objects.create(code="KRE", name="Kaydan Real Estate")
        self.drh = User.objects.create_user("drh.st", password=None, role="DRH", is_group_scope=True)
        self.drh_kre = User.objects.create_user("drh.st.kre", password=None, role="DRH")
        self.drh_kre.subsidiaries.add(self.kre)
        self.lecteur = User.objects.create_user("lect.st", password=None, role="READER", is_group_scope=True)
        source = DataSource.objects.create(name="kshieldApp", slug="kaydan-shield", target_module="rh",
                                           source_type=SourceType.KAYDAN_SHIELD, status=SourceStatus.CONNECTED)
        self.connecteur = DataConnector.objects.create(source=source, base_url="https://api.kaydanshield.com/api/v1")
        patcher = patch("apps.integrations.shield_structure.charge_structure", return_value=_charge())
        patcher.start()
        self.addCleanup(patcher.stop)

    def _get(self, user, query=""):
        self.client.force_authenticate(user)
        return self.client.get(f"/api/v1/integrations/shield/structure/{query}")

    def test_groupe_voit_tout_et_filtre(self):
        d = self._get(self.drh, "?departement=10").json()
        self.assertEqual(d["effectifs"]["total"], 2)

    def test_filiale_sans_correspondance_ne_voit_rien(self):
        d = self._get(self.drh_kre).json()
        self.assertIsNone(d["effectifs"])
        self.assertIn("Hors périmètre", d["restriction"])

    def test_filiale_avec_correspondance_ne_voit_que_sa_societe(self):
        self.connecteur.config = {"company_mapping": {"3": "KRE"}}
        self.connecteur.save()
        d = self._get(self.drh_kre).json()
        self.assertEqual(d["effectifs"]["employes"], 2, "seule la société 3, rattachée à KRE")
        self.assertEqual({r["code"] for r in d["repartitions"]["societe"]}, {"3"})

    def test_droits(self):
        self.assertEqual(self._get(self.lecteur).status_code, 403)
        self.client.force_authenticate(self.drh)
        self.assertEqual(self.client.get("/api/v1/integrations/shield/decouverte/").status_code, 403,
                         "la découverte est réservée aux administrateurs d'intégration")

    def test_export_et_aide_a_la_decision_reprennent_la_structure(self):
        import io

        from openpyxl import load_workbook

        from apps.governance import filtres as F
        from apps.governance.decision_rh import HRDecisionEngine

        self.client.force_authenticate(self.drh)
        r = self.client.get("/api/v1/governance/hr/rapport.xlsx?periode=7j")
        self.assertEqual(r.status_code, 200)
        feuille = load_workbook(io.BytesIO(r.content))["Structure (Shield)"]
        lignes = [l for l in feuille.iter_rows(values_only=True)]
        self.assertIn(("Type de contrat", "CDD", 2, 66.7), lignes)

        with patch("django.utils.timezone.localdate", return_value=AUJ):
            moteur = HRDecisionEngine(self.drh, F.lire({"periode": "7j"}, AUJ))
            constats = {c["id"]: c for c in moteur.constats()}
        self.assertIn("1 contrat(s) arrivent à échéance sous 30 jours", constats["contrats.echeance"]["constat"])
        # Un RH de filiale sans correspondance : ni structure, ni constat tiré du Groupe.
        moteur_kre = HRDecisionEngine(self.drh_kre, F.lire({"periode": "7j"}, AUJ))
        self.assertNotIn("contrats.echeance", {c["id"] for c in moteur_kre.constats()})
